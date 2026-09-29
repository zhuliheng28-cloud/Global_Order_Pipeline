import asyncio
import os
from openpyxl import load_workbook
from playwright.async_api import async_playwright
from login_helpers.dsers_login import wait_or_auto_login_dsers

LOGIN_URL = "https://accounts.dsers.com/accounts/login"

async def run_dsers_rename(excel_path: str, user_data_dir: str, headless: bool = False, progress_callback=None, task_info=None, workers: int = 3, birthday_only: bool = False, update_fields: list = None):
    # 解析要修改的字段
    if update_fields is not None:
        mod_name = any(str(f).lower() in ["name", "姓名", "客户姓名"] for f in update_fields)
        mod_cpf = any(str(f).lower() in ["cpf", "税号", "abn"] for f in update_fields)
        mod_birthday = any(str(f).lower() in ["birthday", "dob", "出生日期", "公司/门店名称", "date of birth"] for f in update_fields)
    elif birthday_only:
        mod_name = False
        mod_cpf = False
        mod_birthday = True
    else:
        mod_name = True
        mod_cpf = True
        mod_birthday = True

    import builtins
    def _print(*args, **kwargs):
        builtins.print(*args, **kwargs)
        if task_info:
            task_info.check_pause()
        if progress_callback:
            progress_callback(" ".join(str(a) for a in args))
            
    print = _print

    print(f"[*] 加载 Excel 文件: {excel_path}")
    # 自动预检：若传入的表格含有被合并订单列，自动进行拆分预处理以保证所有 Shopify 单号作为独立行并发执行
    try:
        from automators.order_template_utils import find_col_by_aliases, clean_order_template_to_script
        check_wb = load_workbook(excel_path, read_only=True)
        first_row = [cell.value for cell in next(check_wb.active.iter_rows(max_row=1))]
        check_wb.close()
        if find_col_by_aliases(first_row, ['被合并订单(订单编号)', '被合并订单', '被合并', '被合并单号', 'merged order', 'merged orders', 'merged']):
            print("[*] 检测到表格包含【被合并订单】列，自动执行 DSers 独立单号拆解预处理...")
            clean_order_template_to_script(excel_path, excel_path, route="B", sw_dsers_rename=True)
    except Exception:
        pass

    try:
        import datetime
        wb = load_workbook(excel_path, data_only=True)
        ws = wb.active

        # 动态识别表头列索引 (1-based)，保持完全向下兼容
        header_map = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=1, column=c).value
            if v:
                header_map[str(v).strip().lower()] = c

        def find_header_col(aliases, exclude=None):
            exclude = [e.lower() for e in exclude] if exclude else []
            for a in aliases:
                a_norm = a.strip().lower()
                if a_norm in header_map:
                    if not any(ex in a_norm for ex in exclude):
                        return header_map[a_norm]
            for a in aliases:
                a_norm = a.strip().lower()
                for h_norm, c_idx in header_map.items():
                    if any(ex in h_norm for ex in exclude):
                        continue
                    if a_norm in h_norm:
                        return c_idx
            return None

        # 1. 网页搜索单号列 (DSers 页面检索主要使用 Shopify 短单号；若无则回退使用 Ak order / 订单编号)
        col_order = find_header_col([
            'shopify order', 'shopify order number', 'shiopify order number', 'shopify单号', 'shopify'
        ], exclude=['aliexpress', 'ak'])
        if not col_order:
            col_order = find_header_col(['ak order', 'ak_order', '订单编号', 'ak单号'], exclude=['shopify', 'shiopify', 'aliexpress']) or 1

        # 备用单号列 (Ak order 与 订单编号 为同一种 18 位长单号，严格与 Shopify 短单号及 AliExpress 平台单号隔离)
        col_ak_order = find_header_col(['ak order', 'ak_order', '订单编号', 'ak单号'], exclude=['shopify', 'shiopify', 'aliexpress'])

        # 2. CPF 列
        col_cpf = find_header_col(['cpf', 'abnnumber', 'abn', '税号', 'cpf(brazil; optional)', 'cpf/abn']) or 3

        # 3. 客户姓名列 (优先 查询结果，其次 客户姓名 / Customer name)
        col_name = find_header_col(['查询结果', 'customer name', 'customer_name', '客户姓名', '姓名', 'contact name', 'contact_person']) or 5

        # 4. 生日列 (对应下单模板的 Date of birth / 公司/门店名称 / 出生日期)
        col_birthday = find_header_col(['date of birth', 'date_of_birth', 'dob', 'birthday', '出生日期', '公司/门店名称', '公司名称', '门店名称'])

        # 5. 状态记录列
        col_status = find_header_col(['执行状态', '状态', '处理状态', 'status'])
        if not col_status:
            col_status = ws.max_column + 1
            ws.cell(row=1, column=col_status, value="执行状态")

        total_count = 0
        unprocessed_rows = []
        for r in range(2, ws.max_row + 1):
            order_val = ws.cell(row=r, column=col_order).value
            if not order_val and col_ak_order:
                order_val = ws.cell(row=r, column=col_ak_order).value
            name_val = str(ws.cell(row=r, column=col_name).value or '').strip() if col_name else ''
            
            if order_val and str(order_val).strip():
                total_count += 1
                st_val = ws.cell(row=r, column=col_status).value
                if not st_val or str(st_val).strip() == "":
                    # 客户姓名中若包含 already change / already changed，直接标记跳过
                    if name_val and ('already' in name_val.lower() and 'change' in name_val.lower()):
                        ws.cell(row=r, column=col_status, value="跳过(已在DSers更改)")
                        continue
                    unprocessed_rows.append(r)

        try:
            wb.save(excel_path)
        except Exception:
            pass

        print(f"[*] 成功载入表格，总单号: {total_count} 条，剩余未处理: {len(unprocessed_rows)} 条")
        print(f"[*] 字段列映射: 单号=列{col_order}, 姓名=列{col_name}, CPF=列{col_cpf}, Birthday=列{col_birthday or '无'}, 状态=列{col_status}")
    except Exception as e:
        print(f"[!] 无法加载 Excel 文件: {e}")
        return

    if not unprocessed_rows:
        print("[*] 表格中所有订单均已处理完毕！无需重复执行。")
        return

    actual_workers = max(1, min(workers, len(unprocessed_rows)))
    print(f"[*] 启动 {actual_workers} 个浏览器窗口并发处理 (待处理 {len(unprocessed_rows)} 条)...")

    # 准备各 Worker 的 session 目录 (平分任务并克隆独立 Session 避免 Chromium 冲突)
    worker_dirs = [user_data_dir]
    for i in range(2, actual_workers + 1):
        wdir = f"{user_data_dir}_worker_{i}"
        if not os.path.exists(wdir) and os.path.exists(user_data_dir):
            import subprocess
            subprocess.run(f"rsync -a --exclude='*Cache*' --exclude='Singleton*' '{user_data_dir}/' '{wdir}/'", shell=True)
        worker_dirs.append(wdir)

    # 切分任务
    worker_row_slices = [unprocessed_rows[i::actual_workers] for i in range(actual_workers)]
    for i in range(actual_workers):
        print(f"[*] Worker-{i+1} 分配到 {len(worker_row_slices[i])} 条订单")

    save_lock = asyncio.Lock()
    processed_counter = 0
    total_unprocessed = len(unprocessed_rows)

    video_dir = os.path.join(os.getcwd(), "videos")
    os.makedirs(video_dir, exist_ok=True)

    async def worker_job(worker_id: int, w_user_data_dir: str, target_rows: list):
        nonlocal processed_counter
        w_tag = f"Worker-{worker_id}"

        for item in ["SingletonLock", "SingletonCookie", "SingletonSocket"]:
            p_lock = os.path.join(w_user_data_dir, item)
            if os.path.exists(p_lock) or os.path.islink(p_lock):
                try: os.remove(p_lock)
                except Exception: pass

        async with async_playwright() as p:
            print(f"[{w_tag}] 启动浏览器 (窗口: {worker_id}, Session: {os.path.basename(w_user_data_dir)})...")
            w_video_dir = os.path.join(video_dir, f"worker_{worker_id}")
            os.makedirs(w_video_dir, exist_ok=True)
            
            context = await p.chromium.launch_persistent_context(
                user_data_dir=w_user_data_dir,
                headless=headless,    
                viewport={'width': 1280, 'height': 800},
                record_video_dir=w_video_dir,
                record_video_size={'width': 1280, 'height': 800}
            )
            if task_info:
                task_info.register_context(context)
            
            page = context.pages[0] if len(context.pages) > 0 else await context.new_page()
            
            # 1. 访问并处理登录
            print(f"[{w_tag}] 访问 {LOGIN_URL} ...")
            await page.goto(LOGIN_URL, wait_until='domcontentloaded', timeout=120000)
            
            await wait_or_auto_login_dsers(page, log_func=print, timeout_seconds=30, task_info=task_info, prefix=f"[{w_tag}]")
                
            print(f"[{w_tag}] 检测到已进入 DSers 首页/主控台！")
            startup_wait = 10 + (worker_id - 1) * 3
            print(f"[{w_tag}] 等待 {startup_wait} 秒缓冲准备进入订单页...")
            for _ in range(startup_wait):
                if task_info:
                    await task_info.async_check_pause()
                await asyncio.sleep(1)

            target_url = "https://www.dsers.com/application/orders/159831080"
            print(f"[{w_tag}] 尝试直接导航至订单页: {target_url}")
            try:
                if target_url not in page.url:
                    await page.goto(target_url, wait_until='domcontentloaded', timeout=60000)
                await asyncio.sleep(5)
            except Exception as e:
                print(f"[{w_tag}] 直接导航失败，尝试点击左侧菜单: {e}")
                try:
                    await page.click("text=Open Orders", timeout=10000)
                    await page.click("text=AliExpress", timeout=10000)
                    await asyncio.sleep(5)
                except Exception as e2:
                    print(f"[{w_tag}] 左侧菜单点击也失败: {e2}")

            search_icon_selector = "svg[data-icon='search'], i.anticon-search, .dsers-icon-search, button[aria-label='Search']"
            try:
                await page.wait_for_selector(search_icon_selector, state="visible", timeout=30000)
            except Exception:
                pass

            # 初始干扰弹窗清理
            try:
                for _ in range(3):
                    closed_count = await page.evaluate('''() => {
                        let closed = 0;
                        const popupContainers = Array.from(document.querySelectorAll(
                            '.ant-popover, .ant-modal, .ant-modal-content, [role="dialog"], .el-dialog, .modal-content, .d-modal, .ant-notification-notice'
                        ));
                        for (const container of popupContainers) {
                            const rect = container.getBoundingClientRect();
                            if (rect.width === 0 || rect.height === 0) continue;
                            if (container.className && typeof container.className === 'string' && container.className.includes('hidden')) continue;
                            const closeEl = container.querySelector(
                                '.anticon-close, [aria-label="close"], [aria-label="Close"], .ant-modal-close, .ant-modal-close-x, [data-icon="close"], button.close, .close'
                            );
                            if (closeEl && (closeEl.offsetParent !== null || closeEl.getBoundingClientRect().width > 0)) {
                                (closeEl.closest('button') || closeEl).click();
                                closed++;
                            }
                        }
                        const directCloseIcons = Array.from(document.querySelectorAll(
                            '.ant-popover .anticon-close, .ant-popover [aria-label="close"], [class*="popoverTitle"] .anticon-close'
                        ));
                        for (const icon of directCloseIcons) {
                            if (icon.offsetParent !== null || icon.getBoundingClientRect().width > 0) {
                                icon.click();
                                closed++;
                            }
                        }
                        return closed;
                    }''')
                    if closed_count > 0:
                        await asyncio.sleep(0.8)
                    else:
                        break
            except Exception:
                pass

            # 遍历分配给本 worker 的行
            for row in target_rows:
                if task_info:
                    await task_info.async_check_pause()

                async with save_lock:
                    st_val = ws.cell(row=row, column=col_status).value
                    if st_val and str(st_val).strip() != "":
                        continue
                    order_id = ws.cell(row=row, column=col_order).value
                    if not order_id and col_ak_order:
                        order_id = ws.cell(row=row, column=col_ak_order).value
                    new_cpf = ws.cell(row=row, column=col_cpf).value
                    new_name = ws.cell(row=row, column=col_name).value
                    new_birthday = ws.cell(row=row, column=col_birthday).value if col_birthday else None

                if not order_id:
                    continue

                order_id = str(order_id).strip()
                new_name = str(new_name).strip() if new_name else ""
                
                # 若名字为 already changed 等标记，直接跳过并回写状态
                if new_name and ('already' in new_name.lower() and 'change' in new_name.lower()):
                    row_status = "跳过(已在DSers更改)"
                    async with save_lock:
                        ws.cell(row=row, column=col_status, value=row_status)
                        try: wb.save(excel_path)
                        except Exception: pass
                        processed_counter += 1
                    continue

                new_cpf = str(new_cpf).strip() if new_cpf and str(new_cpf).strip() != "nan" else ""
                if new_cpf.endswith(".0"):
                    new_cpf = new_cpf[:-2]

                if isinstance(new_birthday, (datetime.date, datetime.datetime)):
                    new_birthday = new_birthday.strftime('%d/%m/%Y')
                new_birthday = str(new_birthday).strip() if new_birthday and str(new_birthday).strip() not in ["nan", "None"] else ""
                if new_birthday.endswith(".0"):
                    new_birthday = new_birthday[:-2]
                if '/' in new_birthday:
                    parts = new_birthday.split('/')
                    if len(parts) == 3 and len(parts[2]) == 4:
                        new_birthday = f"{parts[0].zfill(2)}/{parts[1].zfill(2)}/{parts[2]}"

                active_desc = []
                if mod_name: active_desc.append(f"姓名='{new_name}'")
                if mod_cpf: active_desc.append(f"CPF='{new_cpf}'")
                if mod_birthday: active_desc.append(f"Birthday='{new_birthday}'")
                print(f"\n[{w_tag}][行{row}] 开始处理单号: {order_id} -> 计划修改: {', '.join(active_desc) if active_desc else '未勾选任何修改项'}")


                # --- 步骤 3.1: 搜索 ---
                row_status = "成功"
                try:
                    search_icon_selector = "img.searchSign[src*='search.png']"
                    input_selector = "div.index_flexHeaderItem__GdUPX.searchSign input.ant-input, input.ant-input"
                    
                    try:
                        await page.wait_for_selector(search_icon_selector, state="visible", timeout=5000)
                        await page.click(search_icon_selector, force=True)
                        await page.wait_for_selector(input_selector, state="visible", timeout=5000)
                    except Exception:
                        pass

                    await page.wait_for_selector(input_selector, state="visible", timeout=10000)
                    target_input = page.locator(input_selector).first
                    
                    await target_input.fill("")
                    await target_input.fill(order_id)
                    
                    ok_btn_selector = "button.ant-searchinput-btn"
                    ok_btn = await page.query_selector(ok_btn_selector)
                    if ok_btn and await ok_btn.is_visible():
                        await ok_btn.click(force=True)
                    else:
                        await target_input.press("Enter")
                    
                    try:
                        await page.wait_for_selector(".ant-spin-spinning", state="visible", timeout=1000)
                    except Exception:
                        pass
                    try:
                        await page.wait_for_selector(".ant-spin-spinning", state="hidden", timeout=15000)
                    except Exception:
                        pass
                    await asyncio.sleep(1.0)
                except Exception as e:
                    if task_info and task_info.is_cancel_requested:
                        raise RuntimeError("TASK_CANCELLED_BY_USER")
                    print(f"[{w_tag}] 搜索操作失败: {e}")
                    row_status = "搜索操作失败"
                    async with save_lock:
                        ws.cell(row=row, column=col_status, value=row_status)
                        try: wb.save(excel_path)
                        except Exception: pass
                        processed_counter += 1
                    continue

                if task_info:
                    await task_info.async_check_pause()

                # --- 步骤 3.2: 遍历分类栏寻找订单 ---
                found_category = False
                not_matched_category = None
                try:
                    more_menu_selectors = [".ant-tabs-nav-more", "[aria-label='more']", "span:has-text('...')"]
                    for more_sel in more_menu_selectors:
                        more_btn = await page.query_selector(more_sel)
                        if more_btn and await more_btn.is_visible():
                            await more_btn.hover()
                            await asyncio.sleep(1) 
                            break 

                    tab_selector = ".ant-tabs-tab, [role='tab'], .dsers-tabs-tab, .ant-dropdown-menu-item"
                    tabs = await page.query_selector_all(tab_selector)
                    
                    for tab in tabs:
                        if task_info and task_info.is_cancel_requested:
                            raise RuntimeError("TASK_CANCELLED_BY_USER")
                        if await tab.is_visible():
                            text = await tab.inner_text()
                            if "(1)" in text:
                                lower_text = text.lower()
                                if "failed" in lower_text or "awaiting" in lower_text or "失败" in lower_text or "等待" in lower_text:
                                    print(f"[{w_tag}] 匹配到订单分类: {text.strip()}")
                                    await tab.click()
                                    found_category = True
                                    try:
                                        await page.wait_for_selector("text=Customer Detail", state="visible", timeout=10000)
                                    except Exception:
                                        await asyncio.sleep(1)
                                else:
                                    print(f"[{w_tag}] 订单分类为 {text.strip()} (非 Failed / Awaiting)，跳过。")
                                    not_matched_category = text.strip()
                                break
                except Exception as e:
                    if task_info and task_info.is_cancel_requested:
                        raise RuntimeError("TASK_CANCELLED_BY_USER")
                    print(f"[{w_tag}] 查找分类标签时出错: {e}")

                if not_matched_category:
                    row_status = f"跳过({not_matched_category})"
                    async with save_lock:
                        ws.cell(row=row, column=col_status, value=row_status)
                        try: wb.save(excel_path)
                        except Exception: pass
                        processed_counter += 1
                    continue

                if not found_category:
                    row_status = "link"
                    async with save_lock:
                        ws.cell(row=row, column=col_status, value=row_status)
                        try: wb.save(excel_path)
                        except Exception: pass
                        processed_counter += 1
                    continue

                if task_info:
                    await task_info.async_check_pause()

                # --- 步骤 3.3: 展开订单详情 ---
                try:
                    await page.wait_for_selector("text=Customer Detail", state="visible", timeout=10000)
                    detail_btns = await page.locator("text=Customer Detail").all()
                    clicked = False
                    for btn in detail_btns:
                        if await btn.is_visible():
                            try:
                                await btn.click(timeout=3000)
                            except Exception:
                                await btn.click(force=True)
                            clicked = True
                            break
                    
                    if not clicked:
                        clicked = await page.evaluate('''() => {
                            const elements = Array.from(document.querySelectorAll('*'));
                            for (const el of elements) {
                                if (el.children.length === 0 && el.innerText && el.innerText.trim() === 'Customer Detail') {
                                    const btn = el.closest('button') || el;
                                    btn.click();
                                    return true;
                                }
                            }
                            return false;
                        }''')

                    if clicked:
                        try:
                            await page.wait_for_selector(".ant-drawer-content, .ant-modal-content, [role='dialog']", state="visible", timeout=5000)
                            await page.wait_for_selector(".ant-drawer-content input, .ant-modal-content input, [role='dialog'] input", state="visible", timeout=5000)
                        except Exception:
                            await asyncio.sleep(1)
                    else:
                        raise Exception("Customer Detail 按钮均不可见")
                except Exception as e:
                    print(f"[{w_tag}] 找不到或无法点击 Customer Detail: {e}")
                    row_status = "找不到详情入口"
                    async with save_lock:
                        ws.cell(row=row, column=col_status, value=row_status)
                        try: wb.save(excel_path)
                        except Exception: pass
                        processed_counter += 1
                    continue

                # --- 步骤 3.4: 修改 Contact Name, CPF 与 Birthday ---
                try:
                    await page.evaluate('''() => {
                        let dialogs = Array.from(document.querySelectorAll('.ant-drawer-content, .ant-modal-content, [role="dialog"]')).filter(el => {
                            let rect = el.getBoundingClientRect();
                            return rect.width > 0 && rect.height > 0;
                        });
                        let container = dialogs.length > 0 ? dialogs[dialogs.length - 1] : document.body;
                        let inps = Array.from(container.querySelectorAll('input[type="text"], input:not([type])'));
                        
                        for (let inp of inps) {
                            let rect = inp.getBoundingClientRect();
                            if (rect.width > 0 && rect.height > 0 && !inp.className.includes('search')) {
                                let formItem = inp.closest('.ant-form-item, .ant-row, div[class*="form"], div[class*="item"]');
                                let textContext = formItem ? formItem.innerText : "";
                                if (!textContext) {
                                    let p = inp.parentElement;
                                    let depth = 0;
                                    while (p && p !== container && depth < 6) {
                                        textContext += " " + p.innerText;
                                        p = p.parentElement;
                                        depth++;
                                    }
                                }
                                
                                if (!inp.getAttribute('data-target-input-name') && (textContext.includes("Contact Name") || textContext.includes("Name"))) {
                                    inp.setAttribute('data-target-input-name', 'true');
                                }
                                if (!inp.getAttribute('data-target-input-cpf') && (textContext.includes("Cpf") || textContext.includes("CPF") || textContext.includes("CNPJ") || textContext.includes("Tax ID") || textContext.includes("ID Number"))) {
                                    inp.setAttribute('data-target-input-cpf', 'true');
                                }
                                if (!inp.getAttribute('data-target-input-birthday') && (textContext.toLowerCase().includes("birthday") || textContext.includes("出生日期") || textContext.includes("生日"))) {
                                    inp.setAttribute('data-target-input-birthday', 'true');
                                }
                            }
                        }
                    }''')
                    
                    name_input = page.locator('input[data-target-input-name="true"]').first
                    cpf_input = page.locator('input[data-target-input-cpf="true"]').first
                    birthday_input = page.locator('input[data-target-input-birthday="true"]').first
                    
                    modified_any = False
                    name_found = False
                    birthday_found = False
                    
                    if mod_name:
                        if await name_input.count() > 0:
                            name_found = True
                            current_name = await name_input.input_value()
                            if current_name != new_name and new_name != "":
                                print(f"[{w_tag}] 修改名字: '{current_name}' -> '{new_name}'")
                                await name_input.fill("")
                                await name_input.fill(new_name)
                                modified_any = True
                        else:
                            print(f"[{w_tag}] 找不到 Contact Name 输入框")
                    else:
                        print(f"[{w_tag}] 保持原名字，不执行修改。")
                            
                    if mod_cpf:
                        if await cpf_input.count() > 0 and new_cpf != "" and new_cpf != "nan":
                            current_cpf = await cpf_input.input_value()
                            if current_cpf.strip() != new_cpf:
                                print(f"[{w_tag}] 修改 CPF: '{current_cpf}' -> '{new_cpf}'")
                                await cpf_input.fill("")
                                await cpf_input.fill(new_cpf)
                                modified_any = True
                    else:
                        print(f"[{w_tag}] 保持原CPF，不执行修改。")

                    if mod_birthday:
                        if await birthday_input.count() > 0 and new_birthday != "" and new_birthday != "nan":
                            birthday_found = True
                            current_birthday = await birthday_input.input_value()
                            if current_birthday.strip() != new_birthday.strip():
                                print(f"[{w_tag}] 修改 Birthday: '{current_birthday}' -> '{new_birthday}'")
                                await birthday_input.fill("")
                                await birthday_input.fill(new_birthday.strip())
                                modified_any = True
                            else:
                                print(f"[{w_tag}] Birthday 已匹配 ('{current_birthday}')，无需修改。")
                        elif new_birthday != "" and new_birthday != "nan":
                            print(f"[{w_tag}] 找不到 Birthday 输入框")
                    else:
                        print(f"[{w_tag}] 保持原Birthday，不执行修改。")
                    
                    target_missing = (mod_name and not name_found) or (mod_birthday and not birthday_found)
                    if target_missing:
                        row_status = "找不到对应输入框"
                    elif not modified_any:
                        print(f"[{w_tag}] 目标字段已是最新，无需重复保存。")
                        row_status = "成功(未修改)"
                    else:
                        save_btn = page.locator("button:has-text('Save'), .ant-btn-primary:has-text('Save')").last
                        if await save_btn.is_visible():
                            if await save_btn.is_enabled():
                                await save_btn.click()
                            else:
                                await page.mouse.click(0, 0)
                                await page.wait_for_timeout(500)
                                if await save_btn.is_enabled():
                                    await save_btn.click()
                        
                        try:
                            save_wait1 = asyncio.create_task(page.wait_for_selector(".ant-message-notice-content", state="visible", timeout=4000))
                            save_wait2 = asyncio.create_task(page.wait_for_selector(".ant-drawer-content, .ant-modal-content", state="hidden", timeout=4000))
                            done, pending = await asyncio.wait([save_wait1, save_wait2], return_when=asyncio.FIRST_COMPLETED)
                            for task in pending: task.cancel()
                        except Exception:
                            await asyncio.sleep(1)
                            
                        row_status = "成功"
                    
                    try:
                        await page.evaluate('''() => {
                            document.querySelectorAll('[data-target-input-name], [data-target-input-cpf], [data-target-input-birthday]').forEach(el => {
                                el.removeAttribute('data-target-input-name');
                                el.removeAttribute('data-target-input-cpf');
                                el.removeAttribute('data-target-input-birthday');
                            });
                        }''')
                    except Exception:
                        pass

                except Exception as e:
                    if task_info and task_info.is_cancel_requested:
                        raise RuntimeError("TASK_CANCELLED_BY_USER")
                    print(f"[{w_tag}] 修改过程出错: {e}")
                    row_status = "修改出错"

                async with save_lock:
                    ws.cell(row=row, column=col_status, value=row_status)
                    try:
                        wb.save(excel_path)
                    except PermissionError:
                        print(f"[!] 保存 Excel 失败：表格被占用，请勿在 Excel 软件中打开此表格！")
                    except Exception as e:
                        print(f"[{w_tag}] 保存 Excel 失败: {e}")
                    processed_counter += 1
                    done_now = processed_counter

                try:
                    close_btn = await page.query_selector("button[aria-label='Close']")
                    if close_btn and await close_btn.is_visible():
                        await close_btn.click()
                except Exception:
                    pass

                print(f"[{w_tag}] 进度提示：本次并发已处理 {done_now}/{total_unprocessed} 条 (表格剩余待改: {total_unprocessed - done_now} 条)")
                if task_info:
                    task_info.set_progress(done_now, total_unprocessed, prefix=f"DSers并发改名({actual_workers}路)", unit="条")

            print(f"[{w_tag}] 该 Worker 任务已全部执行完毕！")
            await context.close()

    # 并发执行所有 Workers
    await asyncio.gather(*(
        worker_job(i + 1, worker_dirs[i], worker_row_slices[i])
        for i in range(actual_workers)
    ))

    print("[*] 所有 Worker 任务全部执行完毕！")
    if task_info:
        task_info.set_progress(total_unprocessed, total_unprocessed, prefix="已全部完成", unit="条")

if __name__ == '__main__':
    import argparse
    import sys
    sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    from config import DSERS_SESSION_DIR, SCRIPT_TEMPLATE
    parser = argparse.ArgumentParser(description="DSers 网页端订单自动改名")
    parser.add_argument("--excel", default="/Users/a171325./Documents/landrop/下单9.21.xlsx", help="待处理的 Excel 路径")
    parser.add_argument("--headless", action="store_true", help="是否无头模式")
    parser.add_argument("--workers", type=int, default=3, help="并发浏览器窗口数量")
    parser.add_argument("--birthday-only", action="store_true", default=False, help="仅修改生日，不修改名字和CPF")
    parser.add_argument("--fields", nargs="+", default=None, help="指定修改字段列表，例如: --fields name cpf birthday")
    args = parser.parse_args()
    
    excel_path = args.excel
    if not os.path.exists(excel_path):
        excel_path = SCRIPT_TEMPLATE
    asyncio.run(run_dsers_rename(excel_path=excel_path, user_data_dir=DSERS_SESSION_DIR, headless=args.headless, workers=args.workers, birthday_only=args.birthday_only, update_fields=args.fields))

