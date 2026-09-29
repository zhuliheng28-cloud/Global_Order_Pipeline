import asyncio
import re
import os
import io
from PIL import Image
import ddddocr
from openpyxl import load_workbook
from collections import deque
from playwright.async_api import async_playwright
import sys

# 导入中心化配置
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config import SCRIPT_TEMPLATE, TELEGRAM_SESSION_DIR

# 初始化 OCR 和 DET 目标检测单例，关闭广告输出
ocr = ddddocr.DdddOcr(show_ad=False)
det = ddddocr.DdddOcr(det=True, show_ad=False)

EXCEL_PATH = SCRIPT_TEMPLATE
USER_DATA_DIR = TELEGRAM_SESSION_DIR
CHAT_NAME = "Skynet Robot (Privado)"

def solve_math_captcha(image_bytes):
    """处理并识别数学验证码截图 (支持一位及两位数运算、自适应运算符判定与抗噪)"""
    try:
        # 1. 预处理二值化并检查是否为加载中的模糊缩略图
        image = Image.open(io.BytesIO(image_bytes)).convert('RGB')
        pixels = image.load()
        w, h = image.size
        white_count = 0
        for x in range(w):
            for y in range(h):
                r, g, b = pixels[x, y]
                if r > 160 and g > 160 and b > 160:
                    pixels[x, y] = (255, 255, 255)
                    white_count += 1
                else:
                    pixels[x, y] = (0, 0, 0)
        
        # 正常验证码有效白像素通常在 550~1000 左右，低于 300 说明图片未加载完成（高斯模糊缩略图）
        if white_count < 300:
            print(f"[!] 验证码白像素过少 ({white_count})，判定为未完全加载的缩略图")
            return None, []

        image_L = image.convert('L')
        img_byte_arr = io.BytesIO()
        image_L.save(img_byte_arr, format='PNG')
        clean_bytes = img_byte_arr.getvalue()
        
        # 2. 目标检测
        detected_op = None
        poses = []
        try:
            poses = det.detection(clean_bytes)
            
            # 首先检查 ddddocr 是否直接框出了减号 (宽明显大于高)
            for b in poses:
                bw, bh = b[2] - b[0], b[3] - b[1]
                if bw > 8 and bh > 0 and bw > bh * 2.0 and b[0] < 200:
                    detected_op = '-'
                    break
            
            # 过滤出高度大于 18 的候选数字框 (验证码宽度约为318，等号和问号通常在 x > 180~210)
            digit_boxes = [b for b in poses if (b[3] - b[1] > 18) and b[0] < 210]
            digit_boxes = sorted(digit_boxes, key=lambda b: b[0])
            
            # 如果没直接检测到减号，就去数字中间的缝隙里通过连通域分析找符号
            if not detected_op and len(digit_boxes) >= 2:
                max_gap = -1
                split_idx = 0
                for i in range(len(digit_boxes) - 1):
                    gap = digit_boxes[i+1][0] - digit_boxes[i][2]
                    if gap > max_gap:
                        max_gap = gap
                        split_idx = i
                box1 = digit_boxes[split_idx]
                box2 = digit_boxes[split_idx + 1]
                c1 = (box1[0] + box1[2]) / 2
                c2 = (box2[0] + box2[2]) / 2
                gap_center = (c1 + c2) / 2
                D = c2 - c1
                
                visited = [[False]*h for _ in range(w)]
                components = []
                pix_L = image_L.load()
                
                scan_min_x = max(0, int(box1[0]))
                scan_max_x = min(w, int(box2[2]))
                for x in range(scan_min_x, scan_max_x):
                    for y in range(h):
                        if pix_L[x, y] == 255 and not visited[x][y]:
                            comp_pixels = []
                            q = deque([(x, y)])
                            visited[x][y] = True
                            while q:
                                cx, cy = q.popleft()
                                comp_pixels.append((cx, cy))
                                for dx in [-1, 0, 1]:
                                    for dy in [-1, 0, 1]:
                                        if dx == 0 and dy == 0: continue
                                        nx, ny = cx + dx, cy + dy
                                        if 0 <= nx < w and 0 <= ny < h:
                                            if pix_L[nx, ny] == 255 and not visited[nx][ny]:
                                                visited[nx][ny] = True
                                                q.append((nx, ny))
                            if len(comp_pixels) >= 4:
                                xs = [p[0] for p in comp_pixels]
                                ys = [p[1] for p in comp_pixels]
                                components.append({
                                    'min_x': min(xs), 'max_x': max(xs),
                                    'min_y': min(ys), 'max_y': max(ys),
                                    'cx': (min(xs) + max(xs)) / 2
                                })
                                
                safe_dist = max(5, D / 3.5)
                op_comps = [c for c in components if abs(c['cx'] - gap_center) <= safe_dist]
                if op_comps:
                    min_x = min([c['min_x'] for c in op_comps])
                    max_x = max([c['max_x'] for c in op_comps])
                    min_y = min([c['min_y'] for c in op_comps])
                    max_y = max([c['max_y'] for c in op_comps])
                    op_w = max_x - min_x + 1
                    op_h = max_y - min_y + 1
                    
                    if op_h <= 8 or op_w > op_h * 1.5:
                        detected_op = '-'
                    else:
                        corner_pixels = 0
                        cw = max(1, op_w // 4)
                        ch = max(1, op_h // 4)
                        for y in range(min_y, min_y + ch):
                            for x in range(min_x, min_x + cw):
                                if pix_L[x, y] == 255: corner_pixels += 1
                            for x in range(max_x - cw + 1, max_x + 1):
                                if pix_L[x, y] == 255: corner_pixels += 1
                        for y in range(max_y - ch + 1, max_y + 1):
                            for x in range(min_x, min_x + cw):
                                if pix_L[x, y] == 255: corner_pixels += 1
                            for x in range(max_x - cw + 1, max_x + 1):
                                if pix_L[x, y] == 255: corner_pixels += 1
                        if corner_pixels >= 12:
                            detected_op = '*'
                        else:
                            detected_op = '+'
        except Exception as e:
            print(f"[!] 物理分析运算符出错: {e}")
            
        # 3. OCR 识别全局文字
        res = ocr.classification(clean_bytes)
        print(f"[*] OCR 原始识别结果: {res}")
        res_raw = res.lower().replace('x', '*').replace(' ', '')
        
        # 4. 融合判定运算符
        operator = None
        if '*' in res_raw:
            operator = '*'
        elif detected_op == '-':
            operator = '-'
        elif detected_op == '*':
            operator = '*'
        elif detected_op == '+':
            operator = '+'
        elif '+' in res_raw or '十' in res:
            operator = '+'
        elif '-' in res_raw:
            operator = '-'
            
        if operator:
            print(f"[*] 确认运算符为: {operator}")
            
        def clean_num_str(t):
            replacements = {
                'o': '0', 'O': '0', 'D': '0', 'Q': '0',
                'l': '1', 'I': '1', '|': '1', 'i': '1', '!': '1', 'j': '1', '/': '1', '\\': '1', '[': '1', ']': '1',
                's': '5', 'S': '5',
                'b': '6',
                '乙': '7', '?': '7', '>': '7', '^': '7',
                'B': '8', '&': '8',
                'q': '9', 'g': '9',
                'A': '4', 'u': '4'
            }
            if '7' in res and '2' not in res:
                replacements['z'] = '7'
                replacements['Z'] = '7'
            else:
                replacements['z'] = '2'
                replacements['Z'] = '2'
            for k, v in replacements.items():
                t = t.replace(k, v)
            return t

        # 5. 解析操作数 (支持两位数，如 11 - 2, 14 - 5)
        n1, n2 = None, None
        try:
            if 'digit_boxes' in locals() and len(digit_boxes) >= 2:
                # 寻找最大水平间隙划分左操作数与右操作数
                max_gap = -1
                split_idx = 0
                for i in range(len(digit_boxes) - 1):
                    gap = digit_boxes[i+1][0] - digit_boxes[i][2]
                    if gap > max_gap:
                        max_gap = gap
                        split_idx = i
                        
                left_boxes = digit_boxes[:split_idx+1]
                right_boxes = digit_boxes[split_idx+1:]

                def extract_number_from_boxes(boxes):
                    min_x = min(b[0] for b in boxes)
                    min_y = min(b[1] for b in boxes)
                    max_x = max(b[2] for b in boxes)
                    max_y = max(b[3] for b in boxes)
                    x1, y1 = max(0, min_x - 10), max(0, min_y - 10)
                    x2, y2 = min(w, max_x + 10), min(h, max_y + 10)
                    crop = image_L.crop((x1, y1, x2, y2))
                    arr = io.BytesIO()
                    crop.save(arr, format='PNG')
                    txt = ocr.classification(arr.getvalue())
                    txt = clean_num_str(txt)
                    d = re.sub(r'[^\d]', '', txt)
                    return int(d) if d else None
                    
                n1 = extract_number_from_boxes(left_boxes)
                n2 = extract_number_from_boxes(right_boxes)
                if n1 is not None and n2 is not None:
                    print(f"[*] 物理分块裁剪提取操作数成功: n1={n1}, n2={n2}")
        except Exception as e:
            print(f"[!] 物理裁剪提取数字出错: {e}")

        # 增强：若目标检测分块因连笔杂线未识别，或在减法中出现反常负数 (n1 < n2)，采用自适应左右区域精准重采
        if n1 is None or n2 is None or (operator == '-' and n1 < n2):
            try:
                crop_l = image_L.crop((int(w * 0.05), int(h * 0.1), int(w * 0.35), int(h * 0.9)))
                arr_l = io.BytesIO()
                crop_l.save(arr_l, format='PNG')
                d_l = re.sub(r'[^\d]', '', clean_num_str(ocr.classification(arr_l.getvalue())))

                crop_r = image_L.crop((int(w * 0.43), int(h * 0.1), int(w * 0.65), int(h * 0.9)))
                arr_r = io.BytesIO()
                crop_r.save(arr_r, format='PNG')
                d_r = re.sub(r'[^\d]', '', clean_num_str(ocr.classification(arr_r.getvalue())))

                if d_l and d_r:
                    zn1, zn2 = int(d_l), int(d_r)
                    if operator != '-' or zn1 >= zn2:
                        n1, n2 = zn1, zn2
                        print(f"[*] 自适应区域重采操作数成功: n1={n1}, n2={n2}")
            except Exception as e:
                print(f"[!] 自适应区域重采出错: {e}")
            
        # 如果物理裁剪未完全获取，回退到全局字符串解析
        if n1 is None or n2 is None:
            res_subs = clean_num_str(res)
            nums = re.findall(r'\d+', res_subs)
            if len(nums) >= 2:
                n1, n2 = int(nums[0]), int(nums[-1])
                print(f"[*] 全局字符串提取操作数: n1={n1}, n2={n2}")
            elif len(nums) == 1:
                s = nums[0]
                if len(s) == 2:
                    n1, n2 = int(s[0]), int(s[1])
                    print(f"[*] 单字符串拆分操作数: n1={n1}, n2={n2}")
                elif len(s) >= 3:
                    if operator == '-':
                        n1, n2 = int(s[:-1]), int(s[-1])
                    else:
                        n1, n2 = int(s[0]), int(s[-1])
                    print(f"[*] 智能拆分多位数操作数: n1={n1}, n2={n2}")
        
        ordered_ans = []
        if n1 is not None and n2 is not None:
            ans_plus = n1 + n2
            ans_mul = n1 * n2
            ans_minus = n1 - n2 if n1 >= n2 else None
            
            def add_ans(a):
                if a is not None:
                    ordered_ans.append(a)
            
            if operator == '+':
                add_ans(ans_plus)
                add_ans(ans_minus)
                add_ans(ans_mul)
            elif operator == '-':
                add_ans(ans_minus)
                add_ans(ans_plus)
                add_ans(ans_mul)
            elif operator == '*':
                add_ans(ans_mul)
                add_ans(ans_plus)
                add_ans(ans_minus)
            else:
                add_ans(ans_plus)
                add_ans(ans_minus)
                add_ans(ans_mul)
                    
        certain_ans_str = None
        if n1 is not None and n2 is not None and operator is not None:
            if operator == '+':
                certain_ans_str = str(n1 + n2)
            elif operator == '-':
                if n1 >= n2:
                    certain_ans_str = str(n1 - n2)
            elif operator == '*':
                certain_ans_str = str(n1 * n2)

        # 提取全局 OCR 直接给出的乘法/加减法算式备用答案
        global_direct_ans = None
        gm = re.search(r'(\d+)\s*([\*\+\-])\s*(\d+)', res_raw)
        if gm:
            try:
                gn1 = int(gm.group(1))
                gop = gm.group(2)
                gn2_str = gm.group(3)
                if len(gn2_str) > 1 and gn2_str[-1] in '37j?' and int(gn2_str[:-1]) <= 20:
                    gn2 = int(gn2_str[:-1])
                else:
                    gn2 = int(gn2_str)
                if gop == '*': global_direct_ans = str(gn1 * gn2)
                elif gop == '+': global_direct_ans = str(gn1 + gn2)
                elif gop == '-' and gn1 >= gn2: global_direct_ans = str(gn1 - gn2)
            except Exception:
                pass

        seen = set()
        possible_ans_str = []
        if global_direct_ans:
            possible_ans_str.append(global_direct_ans)
            seen.add(global_direct_ans)
            
        for a in ordered_ans:
            if str(a) not in seen:
                seen.add(str(a))
                possible_ans_str.append(str(a))
                
        # 乘法场景下若物理切块出现异常或全局直解更明确，优先采用全局直解
        if global_direct_ans and (operator == '*' or not certain_ans_str):
            certain_ans_str = global_direct_ans

        if certain_ans_str:
            print(f"[*] 确定的答案为: {certain_ans_str}")
        print(f"[*] 从OCR结果 '{res}' 猜测的可能答案 (按优先级): {possible_ans_str}")
        return certain_ans_str, possible_ans_str
    except Exception as e:
        print(f"[!] 验证码处理出错: {e}")
        return None, []

async def run_cpf_query(excel_path=EXCEL_PATH, user_data_dir=USER_DATA_DIR, headless=False, progress_callback=None, task_info=None):
    def log(msg):
        print(msg)
        if task_info:
            task_info.check_pause()
        if progress_callback:
            progress_callback(msg)
            
    log(f"[*] 加载 Excel 文件: {excel_path}")
    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active
    if ws.max_column < 6 or not ws.cell(row=1, column=6).value:
        ws.cell(row=1, column=6, value="出生日期")
        try:
            wb.save(excel_path)
        except Exception:
            pass

    for item in ["SingletonLock", "SingletonCookie", "SingletonSocket"]:
        p_lock = os.path.join(user_data_dir, item)
        if os.path.exists(p_lock) or os.path.islink(p_lock):
            try: os.remove(p_lock)
            except Exception: pass

    async with async_playwright() as p:
        log("[*] 启动浏览器...")
        context = await p.chromium.launch_persistent_context(
            user_data_dir=user_data_dir,
            headless=headless,
            viewport={'width': 1280, 'height': 800}
        )
        if task_info:
            task_info.register_context(context)
        
        page = context.pages[0] if len(context.pages) > 0 else await context.new_page()
        log(f"[*] 正在打开 Telegram Web...")
        await page.goto('https://web.telegram.org/k/', wait_until='domcontentloaded', timeout=120000)
        
        log(f"[*] 正在左侧聊天列表中寻找 '{CHAT_NAME}' ...")
        try:
            chat_locator = page.locator(f'text="{CHAT_NAME}"').first
            await chat_locator.wait_for(state='visible', timeout=60000)
            await chat_locator.click()
        except Exception as e:
            log(f"[!] 找不到对应的聊天对象，请检查是否在左侧列表中。错误信息: {e}")
            await context.close()
            return
            
        input_selector = 'div.input-message-input'
        await page.wait_for_selector(input_selector, timeout=30000)
        log("[*] 聊天界面加载完毕！开始处理表格数据...")
        
        last_successful_reply_text = ""
        last_processed_cpf = ""
        
        async def process_cpf_row(r_idx, is_retry=False):
            if task_info:
                await task_info.async_check_pause()
            nonlocal last_successful_reply_text, last_processed_cpf
            cpf_cell = ws.cell(row=r_idx, column=4).value # D列
            
            if not cpf_cell:
                if not is_retry:
                    log(f"[*] 第 {r_idx} 行遇到空号码，此遍处理完成！")
                return False
                
            status_cell = ws.cell(row=r_idx, column=5).value # E列
            status_str = str(status_cell).strip() if status_cell else ""
            
            if not is_retry:
                if status_str != "":
                    log(f"[*] 第 {r_idx} 行已有结果 ({status_str})，跳过...")
                    return True
            else:
                if status_str not in ["遇到验证码且未能通过", "查询超时", "提取失败"]:
                    return True

            cpf_text = str(cpf_cell).strip()
            if is_retry:
                log(f"\n[{r_idx}] (重试) 正在重新查询之前失败的 CPF: {cpf_text}")
            else:
                log(f"\n[{r_idx}] 正在查询 CPF: {cpf_text}")
            
            pre_query_text = ""
            pre_in_msg_count = 0
            
            pre_query_msg_id = None
            
            async def close_popup_if_any():
                try:
                    info = await page.evaluate('''() => {
                        // 1. 查找是否存在 Telegram Web K 弹窗层
                        const popup = document.querySelector('.popup.active, .popup, .modal-dialog, .popup-container, div[role="dialog"]');
                        const bodyText = document.body.innerText || "";
                        
                        // 提取弹窗内文本（若找到弹窗容器取容器内文本，否则取 bodyText）
                        let popupText = "";
                        if (popup) {
                            popupText = (popup.innerText || "").trim();
                        } else {
                            let lower = bodyText.toLowerCase();
                            if (lower.includes("resolvido") || lower.includes("já pode continuar") || 
                                lower.includes("errad") || lower.includes("incorret") || lower.includes("tente novamente") ||
                                lower.includes("captcha") || lower.includes("bloque") || lower.includes("aguarde")) {
                                popupText = bodyText.trim();
                            }
                        }
                        
                        // 2. 寻找关闭/确认按钮：优先在弹窗容器内寻找，没有则在全局寻找 OK
                        let targetBtn = null;
                        const searchScope = popup || document;
                        const btns = Array.from(searchScope.querySelectorAll('button, .btn, [role="button"]'));
                        
                        // 寻找特定文字按钮
                        targetBtn = btns.find(b => {
                            const t = (b.innerText || "").trim().toUpperCase();
                            return t === 'OK' || t === 'FECHAR' || t === 'CONFIRMAR' || t === 'ENTENDI' || t === 'CLOSE';
                        });
                        
                        // 若弹窗容器内未找到特定文字，且容器内只有1个主按钮
                        if (!targetBtn && popup) {
                            targetBtn = popup.querySelector('.btn-primary, .popup-button, button');
                        }
                        
                        // 全局保底寻找 OK 按钮 (只要页面上有独立 OK 按钮，即视为弹窗确认按钮)
                        if (!targetBtn) {
                            const allBtns = Array.from(document.querySelectorAll('button, .btn, [role="button"]'));
                            targetBtn = allBtns.find(b => (b.innerText || "").trim().toUpperCase() === 'OK');
                        }
                        
                        if (targetBtn) {
                            const btnLabel = (targetBtn.innerText || "").trim();
                            ['pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click'].forEach(evtType => {
                                targetBtn.dispatchEvent(new MouseEvent(evtType, { bubbles: true, cancelable: true, view: window }));
                            });
                            targetBtn.click();
                            return { closed: true, text: popupText.substring(0, 150), btn: btnLabel };
                        }
                        return { closed: false };
                    }''')
                    if info and info.get('closed'):
                        popup_snippet = info.get('text', '').replace('\n', ' ')
                        log(f"[*] 成功侦测并关闭弹窗 (按钮: '{info.get('btn')}'): {popup_snippet!r}")
                        await asyncio.sleep(0.5)
                        return True
                except Exception as e:
                    log(f"[!] 关闭弹窗时出错: {e}")
                return False

            async def click_btn_robust(btn):
                """高可靠按钮点击：集成视口滚动、全量鼠标/指针事件派发与强制点击，确保 Telegram Web K 响应"""
                try:
                    await btn.scroll_into_view_if_needed()
                    await btn.evaluate('''el => {
                        el.focus();
                        ['pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click'].forEach(evtType => {
                            el.dispatchEvent(new MouseEvent(evtType, { bubbles: true, cancelable: true, view: window }));
                        });
                    }''')
                    await btn.click(force=True, delay=80)
                    return True
                except Exception as e:
                    try:
                        await btn.evaluate('b => b.click()')
                        return True
                    except Exception:
                        return False

            async def scan_and_solve_unresolved_captchas(max_scan=20):
                """向上扫描最近消息 bubbles，寻找未解答的验证码（带图片且带未点击的按钮），解答并点击。若成功处理返回 True"""
                try:
                    recent_bubbles = await page.query_selector_all('div.bubble.is-in')
                    if not recent_bubbles:
                        return False
                    start_idx = len(recent_bubbles) - 1
                    end_idx = max(-1, len(recent_bubbles) - 1 - max_scan)
                    for b_idx in range(start_idx, end_idx, -1):
                        b = recent_bubbles[b_idx]
                        img_el = await b.query_selector('.media-photo')
                        if not img_el:
                            img_el = await b.query_selector('img')
                        btns = await b.query_selector_all('button')
                        if img_el and btns:
                            log(f"[*] 向上扫描发现未解答验证码 (倒数第 {len(recent_bubbles)-b_idx} 条消息)，立即处理...")
                            await asyncio.sleep(0.3)
                            try:
                                # 轮询等待高清验证码加载完成（避免截图到高斯模糊缩略图）
                                certain_ans, possible_answers = None, []
                                for _ in range(8):
                                    image_bytes = await img_el.screenshot()
                                    certain_ans, possible_answers = solve_math_captcha(image_bytes)
                                    if certain_ans or possible_answers:
                                        break
                                    await asyncio.sleep(0.4)
                                    
                                btn_texts = []
                                for btn in btns:
                                    txt = (await btn.inner_text()).strip()
                                    btn_texts.append((btn, txt))
                                target_btn, target_txt = None, None
                                if certain_ans:
                                    for btn, txt in btn_texts:
                                        if txt == certain_ans:
                                            target_btn, target_txt = btn, txt
                                            break
                                if not target_btn and possible_answers:
                                    for ans_str in possible_answers:
                                        for btn, txt in btn_texts:
                                            if txt == ans_str:
                                                target_btn, target_txt = btn, txt
                                                break
                                        if target_btn:
                                            break
                                if target_btn:
                                    log(f"[*] 自动点击遗漏验证码答案 '{target_txt}'...")
                                    await click_btn_robust(target_btn)
                                elif btns:
                                    guessed_btn = btns[0]
                                    log(f"[*] 未得唯一确信解，选择候选答案点击...")
                                    await click_btn_robust(guessed_btn)
                                log("[*] 遗漏验证码已点击，等待并秒关确认弹窗...")
                                for _ in range(15):
                                    if task_info:
                                        await task_info.async_check_pause()
                                    await asyncio.sleep(0.2)
                                    if await close_popup_if_any():
                                        break
                                return True
                            except Exception as e:
                                log(f"[!] 处理遗漏验证码时出错: {e}")
                                return False
                    return False
                except Exception as e:
                    log(f"[!] 扫描未解答验证码异常: {e}")
                    return False

            async def send_query():
                nonlocal pre_query_text, pre_in_msg_count, pre_query_msg_id
                if task_info:
                    await task_info.async_check_pause()
                await close_popup_if_any()
                last_in_msgs = await page.query_selector_all('div.bubble.is-in .message')
                if last_in_msgs:
                    pre_query_text = await last_in_msgs[-1].inner_text()
                    bubble = await last_in_msgs[-1].evaluate_handle('el => el.closest(".bubble")')
                    pre_query_msg_id = await bubble.get_attribute('data-mid') if bubble else None
                else:
                    pre_query_msg_id = None
                pre_in_msg_count = len(last_in_msgs)
                    
                msg_count = len(await page.query_selector_all('div.message'))
                await page.fill(input_selector, cpf_text)
                await page.press(input_selector, 'Enter')
                for _ in range(25):
                    if task_info:
                        await task_info.async_check_pause()
                    await asyncio.sleep(0.2)
                    if len(await page.query_selector_all('div.message')) > msg_count:
                        break

            await send_query()
            
            reply_text = ""
            attempts = 0
            cpf_mismatch_retries = 0
            last_handled_captcha = None
            verification_failed_id = None
            last_handled_leftover_id = None
            is_timeout = False
            
            while True:
                if task_info:
                    await task_info.async_check_pause()
                if attempts >= 100:
                    is_timeout = True
                    break
                await close_popup_if_any()
                await asyncio.sleep(0.2)
                message_elements = await page.query_selector_all('div.bubble.is-in .message')
                if not message_elements:
                    attempts += 1
                    continue
                    
                last_message = message_elements[-1]
                reply_text = await last_message.inner_text()
                
                bubble_el = await last_message.evaluate_handle('el => el.closest(".bubble")')
                msg_id = await bubble_el.get_attribute('data-mid') if bubble_el else None
                
                # 判断是否是新消息：基于 msg_id 能够完全无视虚拟滚动导致的 DOM 数量变化
                if pre_query_msg_id is None:
                    is_new_reply = True
                elif msg_id != pre_query_msg_id:
                    is_new_reply = True
                elif reply_text != pre_query_text:
                    is_new_reply = True
                else:
                    is_new_reply = False
                
                if not is_new_reply:
                    attempts += 1
                    continue
                
                captcha_id = msg_id if msg_id else reply_text
                
                text_lower = reply_text.lower()
                
                # ------ 防残留消息机制 ------
                if reply_text == last_successful_reply_text and cpf_text != last_processed_cpf:
                    if msg_id != last_handled_leftover_id:
                        last_handled_leftover_id = msg_id
                        log(f"[!] 警告：检测到上一个单号 ({last_processed_cpf}) 的延迟残留回复，自动忽略...")
                    attempts += 1
                    continue
                
                negative_keywords = [
                    "não encontrado", "nao encontrado", "inválido", "invalido", 
                    "utilize o comando", "seguido dos", "dígitos do cpf",
                    "sem registro", "não localizado", "nao localizado", "inexistente",
                    "sem dados", "não cadastrado", "nao cadastrado", "não consta", "nao consta",
                    "não existe", "nao existe"
                ]
                if any(k in text_lower for k in negative_keywords):
                    break

                if "nome" in text_lower:
                    # 校验回复中的 CPF 是否与当前查询的 CPF 一致
                    reply_cpf_match = re.search(r'(?i)cpf[\s:]*([0-9\.\-]+)', reply_text)
                    if reply_cpf_match:
                        raw_target_cpf = re.sub(r'\D', '', str(cpf_text))
                        target_cpf = raw_target_cpf[-11:] if len(raw_target_cpf) > 11 else raw_target_cpf
                        raw_reply_cpf = re.sub(r'\D', '', reply_cpf_match.group(1))
                        reply_cpf = raw_reply_cpf[-11:] if len(raw_reply_cpf) > 11 else raw_reply_cpf
                        
                        if target_cpf and reply_cpf and target_cpf != reply_cpf:
                            cpf_mismatch_retries += 1
                            if cpf_mismatch_retries >= 5:
                                log(f"[!] 警告：CPF 连续 {cpf_mismatch_retries} 次不匹配 (期望: {target_cpf}, 实际: {reply_cpf})，达到上限，将结果置空并跳过！")
                                reply_text = ""
                                break
                            log(f"[!] 警告：回复中的 CPF ({reply_cpf}) 与当前查询 ({target_cpf}) 不匹配 (第 {cpf_mismatch_retries}/5 次)！重新发送查询并等待...")
                            await send_query()
                            attempts = 0
                            continue
                    break
                        
                if "muitas requisições" in text_lower or "suspenso" in text_lower or "captcha" in text_lower or "errado" in text_lower:
                    bubble = await last_message.evaluate_handle('el => el.closest(".bubble")')
                    
                    img_el = None
                    img_src = ""
                    if bubble:
                        img_el = await bubble.query_selector('.media-photo')
                        if not img_el:
                            img_el = await bubble.query_selector('img')
                        if img_el:
                            img_src = await img_el.get_attribute('src') or ""
                            
                    captcha_id = f"{msg_id}_{img_src}" if msg_id else f"{reply_text}_{img_src}"
                    
                    if captcha_id == last_handled_captcha:
                        attempts += 1
                        continue
                        
                    log(f"[!] 检测到验证码！开始尝试识别...")
                    
                    await asyncio.sleep(0.5)
                    captcha_id = f"{msg_id}_{img_src}" if msg_id else f"{reply_text}_{img_src}"

                    if not img_el:
                        buttons_check = await bubble.query_selector_all('button') if bubble else []
                        if not buttons_check:
                            if captcha_id != verification_failed_id:
                                verification_failed_id = captcha_id
                                log(f"[!] 消息含关键词但无图片无按钮: {reply_text[:80]!r}")
                                match_min = re.search(r'(\d+)\s*minutos', reply_text.lower())
                                if match_min:
                                    wait_mins = int(match_min.group(1))
                                    log(f"[*] 收到等待提示 ({wait_mins} 分钟)。确认休眠前，先向上深度扫描是否有未解答的验证码...")
                                    solved = await scan_and_solve_unresolved_captchas(max_scan=20)
                                    if solved:
                                        log("[*] 成功处理了未解答的验证码，无需休眠，立即重新发送 CPF 查询...")
                                        send_query.wait_min_attempts = 0
                                        await send_query()
                                        attempts = 0
                                        continue
                                    
                                    # 如果没有找到验证码，且是第一次收到提示，尝试立即重新发送一次刷新
                                    if not getattr(send_query, "wait_min_attempts", False):
                                        send_query.wait_min_attempts = 1
                                        log(f"[*] 未发现待解答验证码，先尝试立即重发一次 CPF 刷新会话...")
                                        await send_query()
                                        await asyncio.sleep(2)
                                        solved_after = await scan_and_solve_unresolved_captchas(max_scan=20)
                                        if solved_after:
                                            log("[*] 重发后扫描解答了新验证码，继续处理...")
                                            send_query.wait_min_attempts = 0
                                            await send_query()
                                            attempts = 0
                                            continue
                                        attempts = 0
                                    else:
                                        log(f"[*] 向上扫描确认无未解答验证码，确实需要等待 {wait_mins} 分钟。开始休眠...")
                                        for _ in range(wait_mins * 60):
                                            if task_info:
                                                await task_info.async_check_pause()
                                            await asyncio.sleep(1)
                                        send_query.wait_min_attempts = 0
                                        log(f"[*] 休眠结束，重新发送 CPF: {cpf_text} ...")
                                        await send_query()
                                        attempts = 0
                                elif "muitas requisições" in reply_text.lower() or "suspenso" in reply_text.lower():
                                    log(f"[*] 判定为频率限制或等待，先扫描屏幕历史消息确认是否有未解答验证码...")
                                    solved = await scan_and_solve_unresolved_captchas(max_scan=20)
                                    if solved:
                                        log("[*] 成功处理了未解答的验证码，无需等待，立即重新发送 CPF 查询...")
                                        await send_query()
                                        attempts = 0
                                        continue
                                    log(f"[*] 无未解答验证码，15 秒后重新发送 CPF: {cpf_text} ...")
                                    for _ in range(15):
                                        if task_info:
                                            await task_info.async_check_pause()
                                        await asyncio.sleep(1)
                                    await send_query()
                                    attempts = 0
                                    log(f"[*] 已重新发送，继续等待回复...")
                                else:
                                    log("[*] 判定为普通警告提示 (如验证码过期)，忽略此消息并继续等待...")
                            else:
                                attempts += 1
                            continue

                    if not bubble:
                        log("[!] 找不到验证码的 bubble 容器")
                        break
                        
                    if not img_el:
                        img_el = last_message
                        
                    target_btn = None
                    target_txt = None
                    possible_answers = None
                    image_bytes = None
                    
                    for scan_attempt in range(4):
                        if scan_attempt > 0:
                            log(f"[!] 第 {scan_attempt} 次尝试未得到确信结果，等待 2 秒后重试扫描...")
                            for _ in range(2):
                                if task_info:
                                    await task_info.async_check_pause()
                                await asyncio.sleep(1)
                            # 重新获取图片元素，以防 DOM 刷新
                            if bubble:
                                img_el = await bubble.query_selector('.media-photo')
                                if not img_el:
                                    img_el = await bubble.query_selector('img')
                            if not img_el:
                                img_el = last_message
                                
                        try:
                            # 轮询等待高清验证码加载完成
                            certain_ans, possible_answers = None, []
                            for _ in range(8):
                                image_bytes = await img_el.screenshot()
                                certain_ans, possible_answers = solve_math_captcha(image_bytes)
                                if certain_ans or possible_answers:
                                    break
                                await asyncio.sleep(0.4)
                            
                            buttons = await bubble.query_selector_all('button')
                            wait_sec = 0
                            while not buttons and wait_sec < 5:
                                if task_info:
                                    await task_info.async_check_pause()
                                log(f"[!] 暂时未找到按钮，等待中 ({wait_sec}s/5s)...")
                                await asyncio.sleep(1)
                                wait_sec += 1
                                buttons = await bubble.query_selector_all('button')
                                
                            log(f"[*] 第 {scan_attempt+1} 次尝试提取到 {len(buttons)} 个按钮")
                            btn_texts = []
                            for btn in buttons:
                                txt = (await btn.inner_text()).strip()
                                btn_texts.append((btn, txt))
                                
                            if certain_ans:
                                for btn, txt in btn_texts:
                                    if txt == certain_ans:
                                        target_btn, target_txt = btn, txt
                                        log(f"[*] 确信答案 '{txt}' 存在于选项中，直接使用。")
                                        break
                                        
                            if not target_btn and possible_answers:
                                for ans_str in possible_answers:
                                    for btn, txt in btn_texts:
                                        if txt == ans_str:
                                            target_btn, target_txt = btn, txt
                                            log(f"[*] 命中高置信度候选答案 '{target_txt}'")
                                            break
                                    if target_btn:
                                        break
                                        
                        except Exception as e:
                            log(f"[!] 获取验证码图片或按钮失败 (可能DOM已刷新): {e}")
                            
                        if target_btn or possible_answers or certain_ans:
                            break
                    
                    if not target_btn and not possible_answers and not image_bytes:
                        attempts += 1
                        await asyncio.sleep(0.5)
                        continue

                    if target_btn:
                        log(f"[*] 正在自动点击确信答案 '{target_txt}'...")
                        try:
                            ok = await click_btn_robust(target_btn)
                            if not ok:
                                raise RuntimeError("click_btn_robust returned False")
                        except Exception as e:
                            log(f"[!] 点击按钮失败 (可能DOM已刷新): {e}")
                            attempts += 1
                            await asyncio.sleep(0.5)
                            continue
                        
                        log("[*] 验证码已解答，正在等待并秒关确认弹窗...")
                        last_handled_captcha = captcha_id
                        for _ in range(15):
                            if task_info:
                                await task_info.async_check_pause()
                            await asyncio.sleep(0.2)
                            if await close_popup_if_any():
                                break
                        await send_query()
                        attempts = 0
                        continue
                        
                    log(f"[!] 无法100%确认答案，或确定的答案不在选项中！")
                    if possible_answers:
                        log(f"[*] 猜测的可能答案: {possible_answers}")
                    else:
                        log(f"[*] 无法猜测出可能答案。")
                        
                    log("[!] 无法自动识别验证码，请人工介入：请手动点击正确答案！(等待60秒)...")
                    manual_solved = False
                    for _ in range(300):
                        if task_info:
                            await task_info.async_check_pause()
                        await asyncio.sleep(0.2)
                        if await close_popup_if_any():
                            manual_solved = True
                            break
                    
                    if manual_solved:
                        log("[*] 检测到人工已解决验证码！重发 CPF...")
                        await send_query()
                        attempts = 0
                        continue
                    else:
                        log("[!] 等待人工超时，根据要求重新扫描并尝试强行盲猜一个答案...")
                        recent_bubbles = await page.query_selector_all('div.bubble.is-in')
                        guessed_btn = None
                        guessed_txt = None
                        for b_idx in range(len(recent_bubbles)-1, max(-1, len(recent_bubbles)-6), -1):
                            gb = recent_bubbles[b_idx]
                            gimg_el = await gb.query_selector('.media-photo')
                            if not gimg_el:
                                gimg_el = await gb.query_selector('img')
                            gbtns = await gb.query_selector_all('button')
                            if gimg_el and gbtns:
                                try:
                                    gimage_bytes = await gimg_el.screenshot()
                                    _, gpos = solve_math_captcha(gimage_bytes)
                                    gbtn_texts = [(btn, (await btn.inner_text()).strip()) for btn in gbtns]
                                    if gpos:
                                        for btn, txt in gbtn_texts:
                                            if txt in gpos:
                                                guessed_btn, guessed_txt = btn, txt
                                                break
                                    if not guessed_btn and gbtn_texts:
                                        guessed_btn, guessed_txt = gbtn_texts[0]
                                except:
                                    if gbtns:
                                        guessed_btn = gbtns[0]
                                break
                                
                        if guessed_btn:
                            log(f"[*] 强行盲猜答案 '{guessed_txt}' 并点击...")
                            await click_btn_robust(guessed_btn)
                            await asyncio.sleep(1)
                            await close_popup_if_any()
                            log("[*] 已盲猜，重发 CPF...")
                            await send_query()
                            attempts = 0
                            continue
                        else:
                            log("[!] 重新扫描未找到验证码，自动测试模式下直接跳过并记录失败。")
                            import time
                            os.makedirs("scratch", exist_ok=True)
                            with open(f"scratch/failed_captcha_{int(time.time())}.png", "wb") as f:
                                f.write(image_bytes)
                            ws.cell(row=r_idx, column=5, value="遇到验证码且未能通过")
                            ws.cell(row=r_idx, column=6, value="")
                            try:
                                wb.save(excel_path)
                            except PermissionError:
                                pass
                            return True # Skip to next row
                

                # 注意：只要进入了 is_new_reply 为 True 的阶段，哪怕它是普通的警告，也属于新消息了
                is_captcha_msg = (
                    "muitas requisições" in text_lower
                    or "suspenso" in text_lower
                    or "captcha" in text_lower
                    or "errado" in text_lower
                )
                if not is_captcha_msg and reply_text.strip() and captcha_id != verification_failed_id:
                    verification_failed_id = captcha_id
                    log(f"[!] 收到非标准回复（无 nome、无验证码）: {reply_text[:80]!r}，跳过等待继续...")
                    break

                attempts += 1
                
            if is_timeout:
                log(f"    -> 等待回复超时。")
                ws.cell(row=r_idx, column=5, value="查询超时")
                ws.cell(row=r_idx, column=6, value="")
            else:
                text_lower = reply_text.lower()
                if "nome" in text_lower:
                    match = re.search(r'(?i)nome[\s:]+([^\n]+)', reply_text)
                    if match:
                        extracted_name = match.group(1).strip()
                        log(f"    -> 成功提取到名字: {extracted_name}")
                        ws.cell(row=r_idx, column=5, value=extracted_name)

                        # 提取出生日期 (例如 NASCIMENTO: 16/10/1979)
                        birth_match = re.search(r'(?i)nascimento[\s:]*([0-9]{2}/[0-9]{2}/[0-9]{4})', reply_text)
                        if not birth_match:
                            birth_match = re.search(r'(?i)nascimento[\s:]*([^\n\r]+)', reply_text)
                        if birth_match:
                            extracted_birth = birth_match.group(1).strip()
                            log(f"    -> 成功提取到出生日期: {extracted_birth}")
                            ws.cell(row=r_idx, column=6, value=extracted_birth)
                        else:
                            ws.cell(row=r_idx, column=6, value="")
                        
                        last_successful_reply_text = reply_text
                        last_processed_cpf = cpf_text
                    else:
                        log(f"    -> 无法用正则提取出名字。原始回复: {reply_text[:50]}...")
                        if is_retry:
                            ws.cell(row=r_idx, column=5, value="")
                        else:
                            ws.cell(row=r_idx, column=5, value="提取失败")
                        ws.cell(row=r_idx, column=6, value="")
                elif any(k in text_lower for k in negative_keywords):
                    log(f"    -> CPF 未找到或无效，跳过此号码。")
                    ws.cell(row=r_idx, column=5, value="无")
                    ws.cell(row=r_idx, column=6, value="")
                else:
                    if cpf_mismatch_retries >= 5:
                        log(f"    -> CPF 多次不匹配，已跳过并置空。")
                        ws.cell(row=r_idx, column=5, value="")
                    elif is_retry:
                        log(f"    -> 重试仍未成功，将结果留空。")
                        ws.cell(row=r_idx, column=5, value="")
                    else:
                        if "muitas requisições" in text_lower:
                            log(f"    -> 验证码处理未能成功通过。")
                            ws.cell(row=r_idx, column=5, value="遇到验证码且未能通过")
                        else:
                            log(f"    -> 未收到包含 nome 的回复。")
                            ws.cell(row=r_idx, column=5, value="提取失败")
                    ws.cell(row=r_idx, column=6, value="")
                
            try:
                wb.save(excel_path)
            except PermissionError:
                log(f"[!] 保存失败：请不要在 Excel 软件中打开此表格文件！会导致文件被锁定。")
                
            await asyncio.sleep(0.2)
            return True

        total_count = 0
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=4).value and str(ws.cell(row=r, column=4).value).strip():
                total_count += 1
        
        log(f"[*] 成功载入表格，待查询 CPF 数量: {total_count} 条")
        log(f"[*] 进度提示：现在是 0/{total_count} (共需查询 {total_count} 条)")
        if task_info:
            task_info.set_progress(0, total_count, prefix="TG查名", unit="条")
        
        processed_count = 0
        row = 2
        while True:
            if task_info:
                await task_info.async_check_pause()
            has_more = await process_cpf_row(row, is_retry=False)
            if not has_more:
                break
            cpf_cell = ws.cell(row=row, column=4).value
            if cpf_cell and str(cpf_cell).strip():
                processed_count += 1
                progress_msg = f"[*] 进度提示：现在是 {processed_count}/{total_count} (共需处理 {total_count} 条，当前已处理 {processed_count} 条)"
                log(progress_msg)
                if task_info:
                    task_info.set_progress(processed_count, total_count, prefix="TG查名", unit="条")
                if progress_callback:
                    progress_callback(progress_msg)
            row += 1

        log("\n=======================================================")
        log("[*] 第一遍查询完成！现在开始重新查询之前失败的号码...")
        log("=======================================================")
        
        retry_rows = []
        for r in range(2, ws.max_row + 1):
            status_val = str(ws.cell(row=r, column=5).value or "").strip()
            if status_val in ["遇到验证码且未能通过", "查询超时", "提取失败"]:
                retry_rows.append(r)
        
        total_retry = len(retry_rows)
        retry_processed = 0
        
        row = 2
        while True:
            if task_info:
                await task_info.async_check_pause()
            has_more = await process_cpf_row(row, is_retry=True)
            if not has_more:
                log(f"[*] 失败重试环节处理完成！")
                break
            if row in retry_rows:
                retry_processed += 1
                progress_msg = f"[*] 重试进度提示：现在是 {retry_processed}/{total_retry} (重试总数: {total_retry} 条，当前重试到第 {retry_processed} 条)"
                log(progress_msg)
                if task_info:
                    task_info.set_progress(retry_processed, total_retry, prefix="TG重试", unit="条")
                if progress_callback:
                    progress_callback(progress_msg)
            row += 1

        log("[*] 所有任务执行完毕！")
        if task_info:
            task_info.set_progress(total_count, total_count, prefix="已完成", unit="条")
        await context.close()

if __name__ == '__main__':
    asyncio.run(run_cpf_query())
